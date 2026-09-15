#!/usr/bin/env python3
"""แปลงไฟล์ 2026 OwnLife Personal Financial Planner (.xlsx) เป็น JSON สำหรับนำเข้าแอป

วิธีใช้:
    python3 tools/convert_xlsx.py path/to/planner.xlsx \
        --out data/import-2026.json --report data/import-report.md

ไฟล์ผลลัพธ์มีข้อมูลการเงินส่วนตัว — อยู่ใน .gitignore และต้องไม่ commit
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
import uuid
from collections import Counter, defaultdict

import openpyxl
from openpyxl.utils import column_index_from_string as ci

# namespace คงที่ เพื่อให้ id เดิมทุกครั้งที่แปลงซ้ำ (ไม่เกิดรายการซ้ำตอน merge)
NS = uuid.UUID("6f1a6d2e-6a2f-5f4b-9c3d-7e5a1b2c3d4e")

SCHEMA_VERSION = 1
APP_ID = "thai-expense-tracker"

# ---------------------------------------------------------------- ผังของชีต
HEADER_ROW = 61
FIRST_DATA_ROW = 62

# บล็อกซ้าย: Income & Saving Tracker / บล็อกขวา: Expense Tracker
BLOCKS = {
    "income": {"date": "B", "type": "F", "group": "J", "amount": "L", "note": "O"},
    "expense": {"date": "W", "type": "Z", "group": "AD", "amount": "AG", "note": "AJ"},
}

# ตารางประเภท (หัวตารางแถว 30, รายการเริ่มแถว 32) + ตารางรายได้ (แถว 13–20)
CATEGORY_TABLES = [
    {"group": "DE", "type_col": "D", "budget_col": "H", "actual_col": "J", "rows": range(32, 48)},
    {"group": "SD", "type_col": "R", "budget_col": "U", "actual_col": "W", "rows": range(32, 48)},
    {"group": "IV", "type_col": "AD", "budget_col": "AH", "actual_col": "AJ", "rows": range(32, 48)},
    {"group": "SV", "type_col": "AQ", "budget_col": "AS", "actual_col": "AT", "rows": range(32, 48)},
    {"group": "IN", "type_col": "AQ", "budget_col": "AS", "actual_col": "AT", "rows": range(13, 21)},
]

# ป้ายกลุ่มในคอลัมน์ Group ของตารางบันทึก
GROUP_BY_LABEL = {
    "รายได้": "IN",
    "รายจ่าย": "DE",
    "ค่าหนี้": "SD",
    "การลงทุน": "IV",
    "การออมเงิน": "SV",
}

GROUP_LABEL_TH = {
    "IN": "รายได้",
    "DE": "รายจ่ายประจำวัน",
    "SD": "ค่ารายเดือน / ค่าหนี้",
    "IV": "การลงทุน",
    "SV": "การออมเงิน",
}

# เซลล์สรุปของชีต ใช้ตรวจทานยอด
TOTALS = {
    "income": "K12",       # =SUM(AT13:AT20)
    "expense_no_sv": "K13",  # =SUM(J31,W31,AJ31)
    "savings": "K14",      # =AT31
    "expense_all": "AR5",  # =sum(J31,W31,AJ31,AT31)
    "remaining": "K16",    # =I10+K12-K13-K14
}
OPENING_BALANCE_COL, OPENING_BALANCE_ROW = "I", 10  # ป้าย "ระบุยอดเริ่มต้น" อยู่ที่ E10, ช่องกรอกคือ I10 (merge I10:L10)

IGNORED_SHEETS = ("Instruction", "Dashboard", "Example")

# เหตุผลที่ "เติมข้อมูลให้" ไม่ใช่ "ข้ามแถว" — แถวพวกนี้ถูกนำเข้าตามปกติ
ADJUSTED_REASONS = (
    "เว้นวันที่ไว้ ใช้วันที่เดียวกับแถวก่อนหน้า",
    "ไม่มีวันที่และไม่มีแถวก่อนหน้า ใช้วันที่ 1 ของเดือน",
)
SCRATCH_SHEET = "Sheet6"


# ------------------------------------------------------------------ utilities
def cell(ws, col: str, row: int):
    return ws.cell(row=row, column=ci(col)).value


def clean_text(v) -> str | None:
    if v is None:
        return None
    s = str(v).replace("​", "").strip()
    if s in ("", "-", "–", "—"):
        return None
    return s


def clean_number(v) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("บาท", "").strip()
    if s in ("", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def clean_date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def round_money(x: float) -> float:
    r = round(x + 0.0, 2)
    return int(r) if r == int(r) else r


def make_id(sheet: str, block: str, row: int) -> str:
    return str(uuid.uuid5(NS, f"{sheet}|{block}|{row}"))


# ------------------------------------------------------------- category tables
def read_category_tables(wb, month_sheets):
    """คืน (categories, month_budgets, table_months)

    categories: dict ชื่อประเภท -> {'name','group'}  (union ของทุกเดือน)
    month_budgets: dict 'YYYY-MM' -> {ชื่อประเภท: งบ}
    table_months: dict group -> list ของเดือนที่ตารางนั้นมีข้อมูล
    """
    categories: dict[str, dict] = {}
    month_budgets: dict[str, dict[str, float]] = defaultdict(dict)
    table_months: dict[str, list[str]] = defaultdict(list)

    for sheet_name, ws, month_key in month_sheets:
        for table in CATEGORY_TABLES:
            found = False
            for row in table["rows"]:
                name = clean_text(cell(ws, table["type_col"], row))
                if not name:
                    continue
                found = True
                # ประเภทเดียวกันอยู่ได้กลุ่มเดียว — เดือนแรกที่พบเป็นเจ้าของ
                categories.setdefault(name, {"name": name, "group": table["group"]})
                budget = clean_number(cell(ws, table["budget_col"], row))
                if budget is not None:
                    month_budgets[month_key][name] = round_money(budget)
            if found:
                table_months[table["group"]].append(month_key)
    return categories, dict(month_budgets), dict(table_months)


# -------------------------------------------------------------- transactions
def read_transactions(ws, sheet_name: str, month_start: dt.date, type_to_group: dict):
    """อ่านทั้งสองบล็อกของชีตหนึ่งเดือน คืน (transactions, skipped, notes)"""
    txs, notes = [], []
    skipped = Counter()
    now_ms = int(dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc).timestamp() * 1000)

    for block, cols in BLOCKS.items():
        last_date = None  # เว้นวันที่ไว้ = วันเดียวกับแถวบน (เจ้าของไฟล์ลงวันที่เฉพาะตอนเปลี่ยนวัน)
        for row in range(FIRST_DATA_ROW, ws.max_row + 1):
            raw_type = clean_text(cell(ws, cols["type"], row))
            raw_note = clean_text(cell(ws, cols["note"], row))
            amount = clean_number(cell(ws, cols["amount"], row))
            date = clean_date(cell(ws, cols["date"], row))

            # แถวว่าง: วันที่เป็นสูตร =F1 อยู่แล้ว จึงดูแค่ประเภท/มูลค่า/โน้ต
            if raw_type is None and raw_note is None and amount is None:
                skipped["แถวว่าง (มีแต่วันที่จากสูตร)"] += 1
                continue

            if date is None:
                if last_date is not None:
                    skipped["เว้นวันที่ไว้ ใช้วันที่เดียวกับแถวก่อนหน้า"] += 1
                    date = last_date
                else:
                    skipped["ไม่มีวันที่และไม่มีแถวก่อนหน้า ใช้วันที่ 1 ของเดือน"] += 1
                    date = month_start
            else:
                last_date = date

            # กลุ่ม: จากคอลัมน์ Group ถ้ามี, ไม่งั้นเดาจากประเภท, ไม่งั้นค่าเริ่มต้นของบล็อก
            label = clean_text(cell(ws, cols["group"], row))
            group = GROUP_BY_LABEL.get(label) if label else None
            group_source = "คอลัมน์ Group" if group else None
            if group is None and raw_type:
                group = type_to_group.get(raw_type)
                group_source = "อนุมานจากประเภท" if group else None
            if group is None:
                group = "IN" if block == "income" else "DE"
                group_source = "ค่าเริ่มต้นของบล็อก"

            kind = "expense" if group != "IN" else "income"

            name = raw_note or raw_type
            if name is None:
                skipped["ไม่มีทั้งโน้ตและประเภท"] += 1
                continue
            name = name[:200]

            txs.append({
                "id": make_id(sheet_name, block, row),
                "kind": kind,
                "name": name,
                "amount": round_money(amount) if amount is not None else None,
                "type": raw_type,
                "group": group,
                "date": date.isoformat(),
                "createdAt": now_ms,
                "updatedAt": now_ms,
                "deleted": False,
            })
            notes.append({
                "sheet": sheet_name,
                "row": row,
                "block": block,
                "group_source": group_source,
                "type": raw_type,
                "amount": amount,
                "date": date,
                "in_month": date.year == month_start.year and date.month == month_start.month,
            })
    return txs, skipped, notes


# --------------------------------------------------------------------- report
def fmt_money(x) -> str:
    if x is None:
        return "—"
    return f"{x:,.2f}".rstrip("0").rstrip(".") if x % 1 else f"{x:,.0f}"


def build_report(months, categories, month_budgets, table_months, opening, scratch_rows, all_txs):
    L = []
    A = L.append
    A("# รายงานการแปลงข้อมูลจากไฟล์ Excel\n")
    A(f"สร้างเมื่อ {dt.datetime.now().strftime('%Y-%m-%d %H:%M')}\n")
    A(f"รวมทั้งหมด **{len(all_txs):,} รายการ** และ **{len(categories)} ประเภท**\n")

    A("\n## 1. ตารางตรวจทานยอด (คำนวณจากรายการ vs. ยอดของชีต)\n")
    A("ยอด *รายจ่าย* ในชีตมีสองตัว: `K13` นับเฉพาะ DE+SD+IV และ `AR5` รวมการออมเงิน (SV) ด้วย")
    A("แอปใช้นิยามเดียวกับ `AR5` คือรายจ่ายรวมการออมเงิน\n")
    A("| เดือน | รายรับ (คำนวณ) | รายรับ (ชีต K12) | ตรง? | รายจ่ายรวมออม (คำนวณ) | รายจ่ายรวมออม (ชีต AR5) | ตรง? |")
    A("|---|---:|---:|:--:|---:|---:|:--:|")

    def ok(a, b):
        return "✅" if abs((a or 0) - (b or 0)) < 0.005 else "❌"

    for m in months:
        A(f"| {m['sheet']} ({m['title']}) | {fmt_money(m['calc_income'])} | {fmt_money(m['sheet_income'])} | "
          f"{ok(m['calc_income'], m['sheet_income'])} | {fmt_money(m['calc_expense_all'])} | "
          f"{fmt_money(m['sheet_expense_all'])} | {ok(m['calc_expense_all'], m['sheet_expense_all'])} |")

    A("\n### รายจ่ายแยกตามกลุ่ม (คำนวณ) เทียบ `K13` ของชีต\n")
    A("| เดือน | DE | SD | IV | รวม DE+SD+IV | ชีต K13 | ตรง? | SV (ออม) | ชีต K14 | ตรง? |")
    A("|---|---:|---:|---:|---:|---:|:--:|---:|---:|:--:|")
    for m in months:
        g = m["calc_groups"]
        sub = g["DE"] + g["SD"] + g["IV"]
        A(f"| {m['sheet']} | {fmt_money(g['DE'])} | {fmt_money(g['SD'])} | {fmt_money(g['IV'])} | {fmt_money(sub)} | "
          f"{fmt_money(m['sheet_expense_no_sv'])} | {ok(sub, m['sheet_expense_no_sv'])} | {fmt_money(g['SV'])} | "
          f"{fmt_money(m['sheet_savings'])} | {ok(g['SV'], m['sheet_savings'])} |")

    A("\n### ทำไมยอดรายจ่ายถึงไม่ตรงในเดือน M2–M9\n")
    A("ยอดรวมในชีตไม่ได้บวกจากรายการโดยตรง แต่บวกผ่านตารางประเภท (`SUMIF` ต่อประเภท)")
    A("รายการที่ไม่ได้กรอกช่อง *ประเภท* จึงไม่ถูกนับเข้ายอดรวมของชีตเลย แม้จะกรอกมูลค่าไว้ครบ\n")
    A("| เดือน | รายการรายจ่ายทั้งหมด | ไม่ได้กรอกประเภท | มูลค่าที่ชีตไม่ได้นับ |")
    A("|---|---:|---:|---:|")
    for m in months:
        gap = round_money(m["calc_expense_all"] - (m["sheet_expense_all"] or 0))
        if m["n_expense"] == 0 and abs(gap) < 0.005:
            continue
        A(f"| {m['sheet']} | {m['n_expense']} | {m['n_no_type']} | {fmt_money(gap)} |")
    A("\nสรุป: **ไม่ใช่ข้อผิดพลาดของตัวแปลง** ตัวเลขที่ตัวแปลงคำนวณได้คือยอดจริงที่เจ้าของไฟล์ใช้จ่ายไป")
    A("ส่วนตัวเลขในชีตคือยอดที่ *ตกหล่น* เพราะสูตรมองไม่เห็นรายการที่ไม่มีประเภท")
    A("ตารางรายประเภทข้างล่างยืนยันว่าเมื่อรายการมีประเภทกำกับ ตัวเลขตรงกันทุกบรรทัด")
    A("หลังนำเข้าแอปแล้ว รายการเหล่านี้จะขึ้นเป็น *ไม่ระบุประเภท* และแก้ประเภทย้อนหลังได้\n")

    A("\n### ตรวจทานรายประเภท (เฉพาะเดือนที่ตารางประเภทในชีตมีตัวเลข)\n")
    A("ตารางนี้พิสูจน์ว่าตัวแปลงอ่าน *มูลค่า* ได้ตรงกับที่ชีตคำนวณ เมื่อรายการนั้นมีประเภทกำกับ\n")
    A("| เดือน | ประเภท | คำนวณ | ชีต (Actual) | ตรง? |")
    A("|---|---|---:|---:|:--:|")
    rows_shown = 0
    for m in months:
        for name, calc, sheet_val in m["per_category"]:
            if abs(calc) < 0.005 and abs(sheet_val or 0) < 0.005:
                continue
            A(f"| {m['sheet']} | {name} | {fmt_money(calc)} | {fmt_money(sheet_val)} | {ok(calc, sheet_val)} |")
            rows_shown += 1
    if rows_shown == 0:
        A("| — | ไม่มีเดือนไหนที่ตารางประเภทมีตัวเลข | | | |")

    A("\n## 2. จำนวนรายการต่อเดือน\n")
    A("| เดือน | รายรับ | รายจ่าย | รวม | ไม่ได้ระบุประเภท | ไม่ได้ระบุมูลค่า | ยอดเริ่มต้น (I10) |")
    A("|---|---:|---:|---:|---:|---:|---:|")
    for m in months:
        A(f"| {m['sheet']} ({m['title']}) | {m['n_income']} | {m['n_expense']} | {m['n_income'] + m['n_expense']} | "
          f"{m['n_no_type']} | {m['n_no_amount']} | {fmt_money(opening.get(m['key'], 0))} |")

    A("\n## 3. แถวที่ข้าม และเหตุผล\n")
    A("| เดือน | เหตุผล | จำนวน |")
    A("|---|---|---:|")
    any_skip = False
    for m in months:
        for reason, n in sorted(m["skipped"].items()):
            if reason in ADJUSTED_REASONS:
                continue
            A(f"| {m['sheet']} | {reason} | {n} |")
            any_skip = True
    if not any_skip:
        A("| — | ไม่มีแถวที่ข้าม | 0 |")
    A("\nแถวว่างคือแถวที่มีแต่วันที่ซึ่งมาจากสูตร `=F1` และไม่มีประเภท มูลค่า หรือโน้ตเลย จึงไม่ใช่รายการจริง")
    A("ส่วน \"ไม่มีทั้งโน้ตและประเภท\" คือแถวที่มีแต่มูลค่าลอย ๆ ไม่มีคำอธิบายว่าคืออะไร\n")

    A("\n### แถวที่นำเข้าแล้วแต่ต้องเติมข้อมูลให้\n")
    A("| เดือน | สิ่งที่เติมให้ | จำนวน |")
    A("|---|---|---:|")
    any_adj = False
    for m in months:
        for reason, n in sorted(m["skipped"].items()):
            if reason not in ADJUSTED_REASONS:
                continue
            A(f"| {m['sheet']} | {reason} | {n} |")
            any_adj = True
    if not any_adj:
        A("| — | ไม่มี | 0 |")
    A("\nในไฟล์เดิม เจ้าของไฟล์ลงวันที่เฉพาะตอนที่เปลี่ยนวัน แถวถัดมาที่เว้นว่างไว้จึงหมายถึง")
    A("วันเดียวกับแถวบน ตัวแปลงเลื่อนวันที่ล่าสุดลงมาให้ตามนั้น (ไม่ใช่ยัดไปวันที่ 1 ของเดือน)\n")

    A("\n## 4. รายการที่วันที่ไม่ตรงกับเดือนของชีต\n")
    out_rows = [(m, r) for m in months for r in m["out_of_month"]]
    if out_rows:
        A("| ชีต | แถว | บล็อก | วันที่ | ประเภท | มูลค่า |")
        A("|---|---:|---|---|---|---:|")
        for m, r in out_rows:
            A(f"| {r['sheet']} | {r['row']} | {r['block']} | {r['date']} | {r['type'] or '—'} | {fmt_money(r['amount'])} |")
        A("\nแอปเก็บรายการตามวันที่จริง รายการเหล่านี้จะไปปรากฏในเดือนที่ถูกต้องเอง")
    else:
        A("ไม่มี — ทุกรายการมีวันที่อยู่ในเดือนของชีตตัวเอง\n")

    A("\n## 5. ประเภทที่ใช้ในรายการแต่ไม่มีในตารางประเภท\n")
    unknown = sorted({t for m in months for t in m["unknown_types"]})
    if unknown:
        A("| ประเภท | จำนวนรายการ | กลุ่มที่แอปกำหนดให้ |")
        A("|---|---:|---|")
        counts = Counter(t["type"] for t in all_txs if t["type"] in set(unknown))
        groups = {}
        for t in all_txs:
            if t["type"] in set(unknown):
                groups.setdefault(t["type"], t["group"])
        for t in unknown:
            A(f"| {t} | {counts[t]} | {groups[t]} ({GROUP_LABEL_TH[groups[t]]}) |")
    else:
        A("ไม่มี — ทุกประเภทที่ใช้มีอยู่ในตารางประเภทของบางเดือน\n")

    A("\n## 6. ตารางประเภทถูกกรอกไว้เดือนไหนบ้าง\n")
    A("| กลุ่ม | เดือนที่มีตาราง |")
    A("|---|---|")
    for g in ["IN", "DE", "SD", "IV", "SV"]:
        ms = table_months.get(g, [])
        A(f"| {g} ({GROUP_LABEL_TH[g]}) | {', '.join(ms) if ms else '— ไม่มีเดือนไหนกรอกเลย'} |")
    A("\nแอปเก็บประเภทไว้เป็นรายการกลางชุดเดียว ไม่ต้องกรอกใหม่ทุกเดือน\n")

    A("\n## 7. งบ / เป้าหมายที่พบ (ใช้เป็นค่าเริ่มต้นในแอป)\n")
    A("| ประเภท | กลุ่ม | งบ/เป้าหมาย |")
    A("|---|---|---:|")
    defaults = {}
    for key in sorted(month_budgets):
        for name, v in month_budgets[key].items():
            defaults.setdefault(name, v)
    for name, meta in sorted(categories.items(), key=lambda kv: (kv[1]["group"], kv[0])):
        A(f"| {name} | {meta['group']} | {fmt_money(defaults.get(name, 0))} |")

    A("\n## 8. Sheet6 (บันทึกกระดาษทด — ไม่นำเข้า)\n")
    if scratch_rows:
        A("```")
        for r in scratch_rows:
            A("  ".join("" if v is None else str(v) for v in r))
        A("```")
        A("\nเป็นการจดเลขไว้ชั่วคราว ไม่มีวันที่หรือประเภท จึงไม่นำเข้า")
    else:
        A("ว่าง\n")
    return "\n".join(L) + "\n"


# ----------------------------------------------------------------------- main
def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--out", default="data/import-2026.json")
    ap.add_argument("--report", default="data/import-report.md")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)

    month_sheets = []
    for i in range(1, 13):
        name = f"M{i}"
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        start = clean_date(cell(ws, "F", 1))
        if start is None:
            print(f"ข้าม {name}: อ่านวันที่ใน F1 ไม่ได้", file=sys.stderr)
            continue
        month_sheets.append((name, ws, f"{start.year:04d}-{start.month:02d}"))

    categories, month_budgets, table_months = read_category_tables(wb, month_sheets)
    type_to_group = {n: m["group"] for n, m in categories.items()}
    known_types = set(categories)

    all_txs, months, opening = [], [], {}

    for sheet_name, ws, month_key in month_sheets:
        start = clean_date(cell(ws, "F", 1))
        txs, skipped, notes = read_transactions(ws, sheet_name, start, type_to_group)
        all_txs.extend(txs)

        ob = clean_number(cell(ws, OPENING_BALANCE_COL, OPENING_BALANCE_ROW))
        if ob:
            opening[month_key] = round_money(ob)

        groups = {g: 0.0 for g in ["IN", "DE", "SD", "IV", "SV"]}
        by_type = defaultdict(float)
        for t in txs:
            groups[t["group"]] += t["amount"] or 0.0
            if t["type"]:
                by_type[t["type"]] += t["amount"] or 0.0
        groups = {g: round_money(v) for g, v in groups.items()}

        per_category = []
        for table in CATEGORY_TABLES:
            for row in table["rows"]:
                name = clean_text(cell(ws, table["type_col"], row))
                if not name:
                    continue
                sheet_actual = clean_number(cell(ws, table["actual_col"], row))
                per_category.append((name, round_money(by_type.get(name, 0.0)), sheet_actual))

        months.append({
            "sheet": sheet_name,
            "key": month_key,
            "title": clean_text(cell(ws, "E", 7)) or month_key,
            "n_income": sum(1 for t in txs if t["kind"] == "income"),
            "n_expense": sum(1 for t in txs if t["kind"] == "expense"),
            "n_no_type": sum(1 for t in txs if t["type"] is None),
            "n_no_amount": sum(1 for t in txs if t["amount"] is None),
            "skipped": skipped,
            "calc_groups": groups,
            "calc_income": groups["IN"],
            "calc_expense_all": round_money(groups["DE"] + groups["SD"] + groups["IV"] + groups["SV"]),
            "sheet_income": clean_number(cell(ws, "K", 12)),
            "sheet_expense_no_sv": clean_number(cell(ws, "K", 13)),
            "sheet_savings": clean_number(cell(ws, "K", 14)),
            "sheet_expense_all": clean_number(cell(ws, "AR", 5)),
            "per_category": per_category,
            "out_of_month": [n for n in notes if not n["in_month"]],
            "unknown_types": {n["type"] for n in notes if n["type"] and n["type"] not in known_types},
        })

    # ประเภทที่ใช้จริงแต่ไม่มีในตาราง -> เพิ่มเข้า list ประเภทกลาง
    for t in all_txs:
        if t["type"] and t["type"] not in categories:
            categories[t["type"]] = {"name": t["type"], "group": t["group"]}

    scratch = []
    if SCRATCH_SHEET in wb.sheetnames:
        sw = wb[SCRATCH_SHEET]
        for row in sw.iter_rows(values_only=True):
            if any(v is not None for v in row):
                scratch.append(row)

    default_budgets = {}
    for key in sorted(month_budgets):
        for name, v in month_budgets[key].items():
            default_budgets.setdefault(name, v)

    # เก็บเป็น override เฉพาะเดือนที่ต่างจากค่าเริ่มต้น
    overrides = {}
    for key, table in month_budgets.items():
        diff = {n: v for n, v in table.items() if default_budgets.get(n) != v}
        if diff:
            overrides[key] = diff

    payload = {
        "app": APP_ID,
        "schemaVersion": SCHEMA_VERSION,
        "exportedAt": dt.datetime.now().astimezone().replace(microsecond=0).isoformat(),
        "settings": {
            "categories": [
                {"name": n, "group": m["group"], "aliases": [], "archived": False}
                for n, m in sorted(categories.items(), key=lambda kv: (kv[1]["group"], kv[0]))
            ],
            "budgets": {"defaults": default_budgets, "months": overrides},
            "openingBalances": opening,
        },
        "transactions": sorted(all_txs, key=lambda t: (t["date"], t["id"])),
    }

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    with open(args.report, "w", encoding="utf-8") as f:
        f.write(build_report(months, categories, month_budgets, table_months, opening, scratch, all_txs))

    print(f"เขียน {args.out}: {len(all_txs):,} รายการ, {len(categories)} ประเภท")
    print(f"เขียน {args.report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
